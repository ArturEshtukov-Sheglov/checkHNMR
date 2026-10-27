"""Structures (SMILES) from compound names with the OPSIN web service, with a molecular-formula check.

Reads the table made by extract_names.py (columns id, name, optionally formula and smiles), asks
OPSIN (https://www.ebi.ac.uk/opsin/) for the SMILES of each name, and, if RDKit is installed,
compares the formula of the SMILES with the expected (HRMS) formula. A SMILES already in the
table (e.g. drawn by hand) is kept. Test results and limitations: README_names_to_smiles.txt.

    python names_to_smiles.py compounds.xlsx [-o smiles.xlsx]

Author: Artur V. Eshtukov-Shcheglov. Code written with the assistance of Claude Opus 5.5
(Anthropic); reviewed and tested by the author. Version 1.1 (October 2026).
"""

import csv
import json
import re
import sys
import time
from urllib.error import HTTPError, URLError
from urllib.parse import quote
from urllib.request import urlopen, Request

try:
    import openpyxl
except ImportError:
    sys.exit("The package openpyxl is missing. Install the packages with:\n    python -m pip install -r requirements.txt\n(or: python -m pip install pymupdf openpyxl rdkit)")
try:
    from rdkit import Chem
    from rdkit.Chem.rdMolDescriptors import CalcMolFormula
    from rdkit import RDLogger
    RDLogger.DisableLog("rdApp.*")
except ImportError:
    Chem = None
    print("RDKit not installed - formulas will not be checked (pip install rdkit).")


# ===== Constants =====

OPSIN = "https://www.ebi.ac.uk/opsin/ws/{}.json"
PAUSE = 0.5                      # seconds between requests - please be polite to the free service
TIMEOUT = 30                     # seconds to wait for one OPSIN answer
MAX_MESSAGE = 150                # characters of an OPSIN error message kept in the status
CSV_SNIFF_BYTES = 2048           # bytes of a CSV file read to detect its delimiter

NEW_COLS = ["smiles", "formula_rdkit", "opsin_status"]
OK_STATUSES = ("OK", "OK (name repaired)", "GIVEN")
COLUMN_WIDTHS = {"name": 60, "description": 80, "hetero": 40, "smiles": 60, "opsin_status": 40}
DEFAULT_WIDTH = 16

# Prefixes after which a hyphen is part of the name, not a line break ("tert-butyl", "bis-")
KEEP_HYPHEN = r"(?:tert|sec|iso|neo|cis|trans|endo|exo|syn|anti|rac|meso|bis|tris|di|tri|tetra)$"


# ===== Input =====

def read_table(path):
    """Returns (headers as written, rows as dicts with lower-case keys)."""
    if path.lower().endswith(".xlsx"):
        ws = openpyxl.load_workbook(path, read_only=True).active
        rows = list(ws.iter_rows(values_only=True))
        head = [str(h).strip() if h is not None else "" for h in rows[0]]
        return head, [dict(zip([h.lower() for h in head], r)) for r in rows[1:] if any(v is not None for v in r)]

    with open(path, encoding="utf-8-sig", newline="") as fh:
        sample = fh.read(CSV_SNIFF_BYTES)
        fh.seek(0)
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        rd = csv.DictReader(fh, dialect=dialect)
        head = [h.strip() for h in rd.fieldnames]
        return head, [{k.strip().lower(): v for k, v in r.items()} for r in rd]


# ===== OPSIN =====

def opsin(name):
    """(SMILES, status) for one name from the OPSIN web service."""
    url = OPSIN.format(quote(name.strip(), safe=""))
    try:
        with urlopen(Request(url, headers={"Accept": "application/json"}), timeout=TIMEOUT) as r:
            data = json.loads(r.read().decode("utf-8"))
    except HTTPError as e:          # OPSIN answers 404 with a JSON message for unparsable names
        try:
            data = json.loads(e.read().decode("utf-8"))
        except Exception:
            return "", "FAILED: HTTP %d" % e.code
    except URLError as e:
        return "", "FAILED: no connection (%s)" % e.reason

    if data.get("status") == "SUCCESS":
        return data.get("smiles", ""), "OK"
    return "", "FAILED: " + str(data.get("message", "not parsed"))[:MAX_MESSAGE]


def repaired(name):
    """Variants of a name without line-break remnants of the PDF text (tried only if OPSIN fails)."""
    v1 = re.sub(r"([\)\]])\s+(?=[a-z])", r"\1", name)            # "...phenyl) propanoate"
    v1 = re.sub(r"-\s+", "-", v1)                                  # "2,2-difluoro- 4-methoxy"
    v1 = re.sub(r"\b(bis|tris|tetrakis)\s+\(", r"\1(", v1)         # "bis (phenylsulfonyl)"
    v2 = re.sub(r"([a-z]{2,})-(?=[a-z]{2})",                        # "cyclooc-ta" (not "tert-butyl")
                lambda m: m.group(0) if re.search(KEEP_HYPHEN, m.group(1)) else m.group(1), v1)
    return [v for v in dict.fromkeys([v1, v2]) if v != name]


def structure_for(name):
    """(SMILES, status): OPSIN on the name, then on its repaired variants if it fails."""
    smi, status = opsin(name) if name else ("", "FAILED: empty name")
    time.sleep(PAUSE)
    if not smi and name and "no connection" not in status:
        for v in repaired(name):
            s2, _ = opsin(v)
            time.sleep(PAUSE)
            if s2:
                return s2, "OK (name repaired)"
    return smi, status


# ===== Formula check =====

def norm_formula(f):
    return re.sub(r"\s+", "", str(f or ""))


def element_counts(f):
    """'C12H5ClF3NO' -> {'C':12,'H':5,'Cl':1,...}; element order and spaces do not matter."""
    counts = {}
    for el, n in re.findall(r"([A-Z][a-z]?)(\d*)", norm_formula(f)):
        counts[el] = counts.get(el, 0) + (int(n) if n else 1)
    return counts


def formula_check(smi, expected, status):
    """(RDKit formula of the SMILES, status with the result of the formula comparison added)."""
    calc = ""
    if smi and Chem is not None:
        m = Chem.MolFromSmiles(smi)
        calc = CalcMolFormula(m) if m else ""
        if m is None:
            status += "; SMILES NOT READ BY RDKIT"
    exp = norm_formula(expected)
    if smi and exp and calc:
        status += "" if element_counts(calc) == element_counts(exp) else "; FORMULA MISMATCH"
    return calc, status


# ===== Run =====

def save_table(rows, cols, out_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "compounds"
    ws.append(cols)
    for r in rows:
        ws.append([r.get(c.lower()) for c in cols])
    for k, c in enumerate(cols):
        ws.column_dimensions[openpyxl.utils.get_column_letter(k + 1)].width = COLUMN_WIDTHS.get(c.lower(), DEFAULT_WIDTH)
    try:
        wb.save(out_path)
    except PermissionError:
        sys.exit(f"Close {out_path} in Excel and run again.")


def main(path, out_path="names_to_smiles_result.xlsx"):
    head, rows = read_table(path)
    keep = [h for h in head if h.lower() not in NEW_COLS]      # all input columns are kept

    n_ok = 0
    for i, r in enumerate(rows, 1):
        name = str(r.get("name") or "").strip()
        given = str(r.get("smiles") or "").strip()
        if given:                                               # a structure already in the table
            smi, status = given, "GIVEN"                        # (e.g. drawn by hand) is kept
        else:
            smi, status = structure_for(name)
        calc, status = formula_check(smi, r.get("formula"), status)

        r["smiles"], r["formula_rdkit"], r["opsin_status"] = smi, calc, status
        n_ok += status in OK_STATUSES
        print(f"{i:4d}/{len(rows)}  {status:20.20s}  {r.get('id', i)}")

    save_table(rows, keep + NEW_COLS, out_path)
    print(f"\nDone: {n_ok} of {len(rows)} with a structure (OK or GIVEN). "
          f"Result: {out_path} - all input columns plus smiles, formula_rdkit, opsin_status "
          f"(rows not marked OK need manual drawing; a SMILES typed into the 'smiles' column is kept).")


def output_name(name):
    """The output file name given with -o; '.xlsx' is added if missing."""
    return name if name.lower().endswith(".xlsx") else name + ".xlsx"


if __name__ == "__main__":
    import argparse
    ap = argparse.ArgumentParser(description="Structures (SMILES) from compound names with the OPSIN web "
                                 "service, checked against the HRMS formulas.")
    ap.add_argument("table", help="table made by extract_names.py (.xlsx or .csv)")
    ap.add_argument("-o", dest="out", default="names_to_smiles_result.xlsx",
                    help="output table (default: names_to_smiles_result.xlsx)")
    a = ap.parse_args()
    main(a.table, output_name(a.out))

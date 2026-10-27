#!/usr/bin/env python3
"""
export_docx.py - export to Word (A4 landscape, one compound per page):
  (1) the Pass-3 review table (sheet "assessment"), or
  (2) an output table of the checking tool, checked_*.xlsx (sheet "checked").
The kind of table is recognised automatically.

Auxiliary script: it only READS an existing review workbook and writes .docx files.
It does not run or change any check; the results are exactly those in the workbook.

Layout: A4 landscape, one compound per page:
    heading (compound id, label, name)
    structure image (from SMILES, RDKit) and the SMILES string
    1H description (and 19F/31P, if present)
    table: item (with a short name of the criterion) | verdict | comment
A first page in every file lists what each item checks (switch off with --no-legend).
    reviewer's mark and comment (if any)
No summary.

Usage:
    python export_docx.py                                    # every .xlsx in data/ -> result/<name>/
    python export_docx.py REVIEW.xlsx -o docx_out            # one file per article (default)
    python export_docx.py REVIEW.xlsx -o docx_out --one-file # everything in one file
    python export_docx.py REVIEW.xlsx --lang en              # English comments and labels
    python export_docx.py REVIEW.xlsx --with-script          # add a column with the script's findings
    python export_docx.py REVIEW.xlsx --articles 1,3,7       # only these articles (by order number)

Needs: openpyxl, python-docx; rdkit (optional - without it no structure images).
"""
import argparse
import io
import re
import sys
from collections import OrderedDict
from pathlib import Path

import importlib.util
import subprocess


def ensure_packages():
    """Check ALL packages before any work: missing ones and broken installations
    (e.g. 'DLL load failed'). Offer to install / repair them (yes/no) instead of failing later."""
    import importlib
    checks = [  # (module to import, pip name, required)
        ("openpyxl", "openpyxl", True),
        ("lxml.etree", "lxml", True),             # used by python-docx
        ("docx", "python-docx", True),
        ("PIL.Image", "pillow", False),           # structure images
        ("rdkit.Chem.Draw", "rdkit", False),      # structure images
    ]
    missing, broken = [], []
    for mod, pip, req in checks:
        try:
            importlib.import_module(mod)
        except ModuleNotFoundError as e:
            if e.name and e.name.split(".")[0] == mod.split(".")[0]:
                missing.append((pip, req))
            else:
                broken.append((pip, req))
        except Exception:                         # ImportError / OSError: installed but broken
            broken.append((pip, req))
    if not missing and not broken:
        return True
    for pip, req in missing:
        print(f"Not installed: {pip}" + ("" if req else "  (optional: structure images)"))
    for pip, req in broken:
        print(f"Installed but broken: {pip}" + ("" if req else "  (optional: structure images)"))
    cmds = []
    if missing:
        cmds.append([sys.executable, "-m", "pip", "install"] + [p for p, _ in missing])
    if broken:
        cmds.append([sys.executable, "-m", "pip", "install", "--force-reinstall", "--no-cache-dir"]
                    + [p for p, _ in broken])
    need_required = any(req for _, req in missing + broken)
    show = "\n".join("    " + " ".join(["python"] + c[1:]) for c in cmds)
    if not sys.stdin.isatty():
        if need_required:
            sys.exit("Fix with:\n" + show)
        print("Continuing without structure images. To add them:\n" + show)
        return False
    answer = input("Install / repair now? [y/n]: ").strip().lower()
    if answer in ("y", "yes", "д", "да"):
        ok = all(subprocess.call(c) == 0 for c in cmds)
        sys.exit(("Done. Run the script again." if ok else "Installation failed. Try manually:\n" + show))
    if need_required:
        sys.exit("Cannot continue without: " + ", ".join(p for p, r in missing + broken if r))
    print("Continuing without structure images.")
    return False


IMAGES_OK = ensure_packages()

import openpyxl
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

try:
    from rdkit import Chem
    from rdkit.Chem import Draw
    HAVE_RDKIT = True
except ImportError:
    HAVE_RDKIT = False

# ---------------------------------------------------------------- columns
# Columns are found by their header text (first row of the sheet), not by position.
COL = {
    "order": "Порядок",
    "id": "соединение",
    "label": "метка",
    "name": "название (SI)",
    "smiles": "использованный SMILES",
    "source": "источник структуры",
    "desc": "описание(я) 1H",
    "xnuc": "19F/31P",
    "reasons_ru": "причины (Claude)",
    "reasons_en": "reasons (Claude), EN original",
    "agree": "Артур: согласен?",
    "comment": "Артур: комментарий",
    "script_ru": "скрипт (замороженная версия 18.10): нахо",   # prefix match
    "script_en": "script (frozen 18.10): findings, EN orig",    # prefix match
}
ITEMS = [  # (item, header of the verdict column)
    ("1", "пункт 1"),
    ("2", "пункт 2"),
    ("3", "пункт 3"),
    ("4", "пункт 4"),
    ("6", "пункт 6"),
    ("7", "пункт 7"),
]

# ---------------------------------------------------------------- wording
VERDICT_RU = {
    "consistent": "согласовано",
    "inconsistent": "несогласовано",
    "partner unverifiable": "партнёр непроверяем",
    "cannot be determined": "определить нельзя",
    "mixture/rotamers": "смесь/ротамеры",
    "subcategory: exchangeable proton not reported": "подкатегория: обменный протон не указан",
    "subcategory: protons on quadrupolar nuclei not reported": "подкатегория: протоны при квадрупольных ядрах не указаны",
    "subcategory: flexible chain": "подкатегория: гибкая цепь",
    "subcategory: borderline strong coupling": "подкатегория: пограничная сильная связь",
    "subcategory: broadening without app": "подкатегория: уширение без app",
}
AGREE_RU = {"yes": "да", "no": "нет", "partly": "частично"}
TXT = {
    "ru": dict(item="Пункт", verdict="Вердикт", comment="Комментарий", script="Скрипт",
               desc="Описание", assign="Отнесения", reviewer="Проверяющий",
               article="Статья", source="структура", other="Прочее",
               contents="Содержание", has_struct="структура есть", no_struct="структуры нет"),
    "en": dict(item="Item", verdict="Verdict", comment="Comment", script="Script",
               desc="Description", assign="Assignments", reviewer="Reviewer",
               article="Article", source="structure", other="Other",
               contents="Contents", has_struct="structure", no_struct="no structure"),
}
# Short names (table) and descriptions (legend page) of the checklist items,
# following the Pass-3 protocol v1.2 (items 5 and 8 are not applied).
CRIT = {
    "ru": {
        "1": ("Сигналы и число протонов",
              "Число протонов в описании совпадает с числом необменивающихся протонов структуры; "
              "отсутствующий сигнал допустим, если причина указана."),
        "2": ("Химические сдвиги",
              "Сдвиги согласуются с характеристическими диапазонами фрагментов структуры (Pretsch, 2020)."),
        "3": ("Партнёры по КССВ",
              "У каждой приведённой КССВ есть сигнал-партнёр с той же КССВ (допуск 0.3 Гц или шаг записи)."),
        "4": ("Мультиплетность",
              "Число КССВ соответствует мультиплетности; мультиплетность не требует больше партнёров, "
              "чем даёт структура."),
        "6": ("Нотация первого порядка",
              "Системы не первого порядка (магнитно неэквивалентные ядра, Δν/J < 6) не записаны как "
              "сигналы первого порядка с КССВ."),
        "7": ("Кажущаяся мультиплетность",
              "Мультиплетность, возникающая лишь при случайном равенстве КССВ с неэквивалентными соседями, "
              "помечена как кажущаяся (app, J_app) или приведена с отдельными КССВ."),
        "note": "Пункты 5 и 8 чек-листа не применяются. Подкатегории, «партнёр непроверяем» и "
                "«определить нельзя» несогласованностью не считаются.",
        "title": "Проверяемые критерии",
        "colours": "Цвет ячейки «Вердикт»: красный – несогласовано; жёлтый – подкатегория, партнёр "
                   "непроверяем или определить нельзя; зелёный – согласовано.",
    },
    "en": {
        "1": ("Signals and proton count",
              "The number of protons reported matches the non-exchangeable protons of the structure; "
              "a missing signal is acceptable if the reason is stated."),
        "2": ("Chemical shifts",
              "Shifts agree with the characteristic ranges of the structural fragments (Pretsch, 2020)."),
        "3": ("Coupling partners",
              "Every reported J has a partner signal with the same J (tolerance 0.3 Hz or the reporting step)."),
        "4": ("Multiplicity",
              "The number of J values matches the multiplicity; the multiplicity needs no more coupling "
              "partners than the structure provides."),
        "6": ("First-order notation",
              "Non-first-order systems (magnetically non-equivalent nuclei, Δν/J < 6) are not reported "
              "as first-order signals with J values."),
        "7": ("Apparent multiplicity",
              "A multiplicity that arises only from accidentally equal couplings to non-equivalent neighbours "
              "is marked as apparent (app, J_app) or reported with separate J values."),
        "note": "Checklist items 5 and 8 are not applied. Subcategories, 'partner unverifiable' and "
                "'cannot be determined' are not counted as inconsistencies.",
        "title": "Criteria checked",
        "colours": "Verdict cell colour: red - inconsistent; yellow - subcategory, partner unverifiable "
                   "or cannot be determined; green - consistent.",
    },
}

# Cell shading as in the workbook: red = inconsistent, yellow = subcategory /
# partner unverifiable / cannot be determined, green = consistent.
SHADE = {"inconsistent": "F4CCCC", "consistent": "D9EAD3"}
SHADE_OTHER = "FFF2CC"


# ---------------------------------------------------------------- helpers
def find_col(header, name):
    """Index of the column whose header equals name (or starts with it)."""
    for i, h in enumerate(header):
        if h is not None and str(h).strip() == name:
            return i
    for i, h in enumerate(header):
        if h is not None and str(h).strip().startswith(name):
            return i
    return None


# A comment segment starts at a sentence beginning with an item marker:
# "Пункт 3: ...", "Item 3: ...", "7: ...", "6(b) borderline ...".
MARK = r"(?:(?:Пункт|Item)\s+)?([1-7])(?:\([a-c]\))?(?=[:\s,])"
SPLIT = re.compile(r"(?<=[.;])\s+(?=" + MARK.replace("([1-7])", "[1-7]")
                   + r"|Отнесения|Assignments|[A-ZА-ЯЁ])")
START = re.compile(r"^" + MARK)
ASSIGN = re.compile(r"^(Отнесения|Assignments)")


def split_reasons(text):
    """Split the free-text reasons into {item: text}; 'assign' for inferred assignments."""
    out = OrderedDict()
    if not text:
        return out
    current = "other"
    for sent in SPLIT.split(str(text).strip()):
        sent = sent.strip()
        if not sent:
            continue
        m = START.match(sent)
        if m:
            current = m.group(1)
        elif ASSIGN.match(sent):
            current = "assign"
        out.setdefault(current, []).append(sent)
    return OrderedDict((k, " ".join(v)) for k, v in out.items())


def split_script(text):
    """Script findings are line-based: 'п. 3: ...' / 'item 3: ...' plus continuation lines."""
    out = OrderedDict()
    if not text:
        return out
    current = "other"
    for line in str(text).splitlines():
        m = re.match(r"^\s*(?:п\.|item)\s*([1-7])", line, re.I)
        if m:
            current = m.group(1)
        out.setdefault(current, []).append(line.strip())
    return OrderedDict((k, "\n".join(v)) for k, v in out.items())


def shade(cell, rgb):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), rgb)
    tcPr.append(shd)


_DRAW_FAILED = False


def structure_png(smiles):
    """PNG of the structure, or None. A drawing failure disables images once, with a warning."""
    global _DRAW_FAILED
    if not HAVE_RDKIT or not IMAGES_OK or _DRAW_FAILED or not smiles:
        return None
    try:
        return _structure_png(smiles)
    except Exception as e:
        _DRAW_FAILED = True
        print(f"Structure images switched off: drawing failed ({type(e).__name__}: {e})")
        return None


def _structure_png(smiles):
    mol = Chem.MolFromSmiles(str(smiles))
    if mol is None:
        return None
    img = Draw.MolToImage(mol, size=(600, 380))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)
    return buf


# English wording of the "structure source" column (the review workbook has it in Russian)
SOURCE_EN = [
    ("OPSIN (название), формула проверена", "OPSIN (name), formula checked"),
    ("OPSIN (название), HRMS для проверки нет", "OPSIN (name), no HRMS to check against"),
    ("нарисованная структура повторяющегося звена", "drawn structure of the repeat unit"),
    ("нарисованная структура", "drawn structure"), ("изображённая структура", "drawn structure"),
    ("SI (агент)", "SI (agent)"), ("рисунка в SI нет", "no drawing in the SI"),
    ("название", "name"), ("по образцу", "after"), ("с. ", "p. "),
]


def source_text(src, lang):
    """The structure-source note; in English, Russian phrases are translated, and a note that is
    still partly Russian is shortened to its first (translated) part."""
    src = str(src)
    if lang != "en":
        return src
    for ru, en in SOURCE_EN:
        src = src.replace(ru, en)
    if re.search("[А-Яа-яЁё]", src):
        src = re.split(r"[,(]", src)[0].strip() if not re.search("[А-Яа-яЁё]", re.split(r"[,(]", src)[0]) \
            else "structure from the SI"
    return src


def greek(text):
    """The tool writes 'dnu' (plain ASCII in Excel); in Word it is shown as Δν."""
    return re.sub(r"\bdnu\b", "Δν", text)


def safe(s):
    return re.sub(r"[^A-Za-z0-9._-]+", "_", str(s)).strip("_")


# ---------------------------------------------------------------- document
def new_document():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(11)
    st.paragraph_format.space_after = Pt(2)
    # headings: same font as the text, black (Word default is blue Calibri Light)
    for name, size in (("Heading 1", 14), ("Heading 2", 12), ("Title", 16)):
        h = doc.styles[name]
        h.font.name = "Times New Roman"
        h.font.size = Pt(size)
        h.font.bold = True
        h.font.color.rgb = RGBColor(0, 0, 0)
        rpr = h.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is not None:
            for a in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
                rfonts.attrib.pop(qn("w:" + a), None)
            for a in ("ascii", "hAnsi", "eastAsia", "cs"):
                rfonts.set(qn("w:" + a), "Times New Roman")
    for sec in doc.sections:                     # A4 landscape
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
        sec.left_margin = sec.right_margin = Cm(2)
        sec.top_margin = sec.bottom_margin = Cm(1.5)
    return doc


# ---------------------------------------------------------------- signals in the description
# One signal of a description, as the checking tool reads it: "7.41 (d, J = 6.9 Hz, 2H)" or a range
# "7.35-7.28 (m, 2H)" (the same pattern as description_parser.SIGNAL of the tool)
SIG_NUM = r"-?\d+(?:\.\d+)?"
SIGNAL = re.compile(r"(?<![\d.])(" + SIG_NUM + r")(?:\s?[–—−-]\s?(" + SIG_NUM + r"))?\s?(?:,\s?(?=\())?(?:ppm\s?)?"
                    r"\(([^()]*(?:\([^()]*\)[^()]*)*)\)")
SIG_REF = re.compile(r"#(\d+)(?: (-?\d+\.\d{2}))?")
FINDING_COLS = [  # (findings header start, verdict header start)
    ("findings 3", "item 3\npartners"), ("findings 4(a)", "item 4(a)"), ("findings 4(b)", "item 4(b)"),
    ("findings 6(a)", "6(a)"), ("findings 6(b)", "6(b)")]
SIG_RED, SIG_YELLOW = "F8CBAD", "FFE699"


def signal_spans(desc):
    """[(start, end, shift as the tool prints it)] of the signals, numbered as the tool numbers them."""
    i = desc.find("δ")
    off = i if i >= 0 else 0
    spans, seen = [], set()
    for m in SIGNAL.finditer(desc[off:]):
        key = re.sub(r"\s", "", m.group(0))
        if key in seen:
            continue
        seen.add(key)
        try:
            a = float(m.group(1))
            sh = (a + float(m.group(2))) / 2 if m.group(2) else a
        except ValueError:
            sh = None
        spans.append((off + m.start(), off + m.end(), None if sh is None else f"{sh:.2f}"))
    return spans


def is_inconsistent(v):
    v = (v or "").lower()
    return v.startswith("inconsistent") and "?" not in v


def signal_marks(row, col):
    """{signal number: colour} from the findings of the tool: red - inconsistent, yellow - to check."""
    marks = {}
    for fh, vh in FINDING_COLS:
        fi, vi = col(fh), col(vh)
        if fi is None or row[fi] is None:
            continue
        verdict = "" if vi is None or row[vi] is None else str(row[vi])
        for line in str(row[fi]).split("\n"):
            refs = SIG_REF.findall(line)
            if not refs or "not evaluated" in line:
                continue
            red = ("INCONSISTENT" in line or "singlets" in line or "impossible" in line
                   or ("likely item-3 inconsistency" in line and is_inconsistent(verdict))
                   or (fh in ("findings 4(a)", "findings 6(a)") and is_inconsistent(verdict)))
            for k, sh in refs:
                k = int(k)
                if red:
                    marks[(k, sh)] = SIG_RED
                elif marks.get((k, sh)) != SIG_RED:
                    marks[(k, sh)] = SIG_YELLOW
    return marks


def add_description(p, desc, marks):
    """The description with the signals coloured as the findings say."""
    spans = signal_spans(desc)
    colour = {}
    for (k, sh), c in marks.items():
        j = k - 1
        if not (0 <= j < len(spans) and (sh is None or spans[j][2] == sh)):
            j = next((n for n, x in enumerate(spans) if sh is not None and x[2] == sh), None)
        if j is not None and colour.get(j) != SIG_RED:
            colour[j] = c
    pos = 0
    for j, (a, b, _) in enumerate(spans):
        if j not in colour:
            continue
        if a > pos:
            p.add_run(desc[pos:a])
        r = p.add_run(desc[a:b])
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), colour[j])
        r._r.get_or_add_rPr().append(shd)
        pos = b
    p.add_run(desc[pos:])


# ---------------------------------------------------------------- contents
def struct_mark(smiles, lang):
    """'(structure)' / '(no structure)': whether the compound has a structure (SMILES) to draw."""
    t = TXT[lang]
    return f"({t['has_struct'] if smiles else t['no_struct']})"


def add_bookmark(paragraph, name, bid):
    """A bookmark around the paragraph, the target of a link in the contents."""
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bid))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bid))
    paragraph._p.insert(0, start) if paragraph._p.pPr is None else paragraph._p.pPr.addnext(start)
    paragraph._p.append(end)


TOC_LINES_PER_PAGE = 36          # lines of the contents on one landscape A4 page (estimate)
TEXT_WIDTH = Cm(25.7)            # landscape A4 minus margins: right tab stop of the page numbers


def add_field(paragraph, instr, shown):
    """A Word field (e.g. PAGEREF) with a provisional result; Word updates it (F9 / on opening)."""
    def fld(kind):
        r = OxmlElement("w:r")
        f = OxmlElement("w:fldChar")
        f.set(qn("w:fldCharType"), kind)
        r.append(f)
        return r
    paragraph._p.append(fld("begin"))
    r = OxmlElement("w:r")
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    r.append(it)
    paragraph._p.append(r)
    paragraph._p.append(fld("separate"))
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = str(shown)
    r.append(t)
    paragraph._p.append(r)
    paragraph._p.append(fld("end"))


def update_fields_on_open(doc):
    """Word offers to update the fields (page numbers of the contents) when the file is opened."""
    st = doc.settings.element
    if st.find(qn("w:updateFields")) is None:
        u = OxmlElement("w:updateFields")
        u.set(qn("w:val"), "true")
        st.append(u)


def add_contents(doc, entries, lang, legend=True):
    """Contents after the criteria page: one line per compound with dot leaders and its page number;
    the line is a link to the compound (Ctrl+click in Word; also a link in a PDF saved from Word).
    Page numbers are PAGEREF fields: estimated here (one compound per page), exact after Word updates."""
    from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
    doc.add_heading(TXT[lang]["contents"], level=1)
    first = (2 if legend else 1) + -(-(len(entries) + 2) // TOC_LINES_PER_PAGE)
    for k, (text, anchor) in enumerate(entries):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.tab_stops.add_tab_stop(TEXT_WIDTH, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        link = OxmlElement("w:hyperlink")
        link.set(qn("w:anchor"), anchor)
        link.set(qn("w:history"), "1")
        r = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        col = OxmlElement("w:color")
        col.set(qn("w:val"), "1F4E9A")
        rpr.append(col)
        r.append(rpr)
        tx = OxmlElement("w:t")
        tx.text = text
        tx.set(qn("xml:space"), "preserve")
        r.append(tx)
        link.append(r)
        p._p.append(link)
        p.add_run("\t")
        add_field(p, f"PAGEREF {anchor} \\h", first + k)
    update_fields_on_open(doc)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def anchor_name(i):
    return f"cpd_{i:05d}"


def add_legend(doc, lang):
    """First page: what each item checks and what the colours mean."""
    cr = CRIT[lang]
    doc.add_heading(cr["title"], level=1)
    tab = doc.add_table(rows=0, cols=2)
    tab.style = "Table Grid"
    tab.autofit = False
    for item, _ in ITEMS:
        cells = tab.add_row().cells
        cells[0].text = f"{item}. {cr[item][0]}"
        cells[0].paragraphs[0].runs[0].bold = True
        cells[1].text = cr[item][1]
    for i, w in enumerate((Cm(5.5), Cm(20.2))):
        tab.columns[i].width = w
        for r in tab.rows:
            r.cells[i].width = w
    doc.add_paragraph()
    doc.add_paragraph(cr["note"])
    doc.add_paragraph(cr["colours"])
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def review_head(row, c, lang):
    get = lambda key: row[c[key]] if c.get(key) is not None else None
    head = f"{get('id')}"
    if get("label"):
        head += f" ({get('label')})"
    return head + " " + struct_mark(get("smiles"), lang)


def add_compound(doc, row, c, lang, with_script, new_page, anchor=None):
    """One compound per page: a page break before every compound except the first."""
    t = TXT[lang]
    get = lambda key: row[c[key]] if c.get(key) is not None else None
    if new_page:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    head = review_head(row, c, lang)
    h = doc.add_heading(head, level=2)
    if anchor:
        add_bookmark(h, anchor, int(anchor[4:]))
    if get("name"):
        doc.add_paragraph(str(get("name"))).runs[0].italic = True

    png = structure_png(get("smiles"))
    if png is not None:
        doc.add_picture(png, width=Cm(9))
    if get("smiles"):
        p = doc.add_paragraph()
        r = p.add_run(str(get("smiles")))
        r.font.size = Pt(8)
        r.font.name = "Consolas"
        if get("source"):
            r2 = p.add_run(f"   ({t['source']}: {source_text(get('source'), lang)})")
            r2.font.size = Pt(8)

    p = doc.add_paragraph()
    p.add_run(f"{t['desc']}: ").bold = True
    p.add_run(str(get("desc") or ""))
    if get("xnuc"):
        doc.add_paragraph(str(get("xnuc")))

    reasons = split_reasons(get("reasons_ru" if lang == "ru" else "reasons_en"))
    script = split_script(get("script_ru" if lang == "ru" else "script_en")) if with_script else {}

    ncol = 4 if with_script else 3
    tab = doc.add_table(rows=1, cols=ncol)
    tab.style = "Table Grid"
    tab.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = [t["item"], t["verdict"], t["comment"]] + ([t["script"]] if with_script else [])
    for i, h in enumerate(hdr):
        tab.rows[0].cells[i].text = h
        tab.rows[0].cells[i].paragraphs[0].runs[0].bold = True

    for item, _ in ITEMS:
        v = row[c["v" + item]] if c.get("v" + item) is not None else None
        v = "" if v is None else str(v)
        cells = tab.add_row().cells
        cells[0].text = f"{item}. {CRIT[lang][item][0]}"
        cells[1].text = VERDICT_RU.get(v, v) if lang == "ru" else v
        cells[2].text = reasons.get(item, "")
        if with_script:
            cells[3].text = script.get(item, "")
        if v:
            shade(cells[1], SHADE.get(v, SHADE_OTHER))
    for key, label in (("assign", t["assign"]), ("other", t["other"])):
        if reasons.get(key):
            cells = tab.add_row().cells
            cells[0].text = label
            cells[1].merge(cells[ncol - 1]).text = reasons[key]

    widths = [Cm(4.6), Cm(4.0), Cm(17.1)] if not with_script else [Cm(4.2), Cm(3.4), Cm(10), Cm(8.1)]
    tab.autofit = False
    for i, w in enumerate(widths):
        tab.columns[i].width = w
    for r in tab.rows:
        for i, w in enumerate(widths):
            r.cells[i].width = w
            for par in r.cells[i].paragraphs:
                for run in par.runs:
                    run.font.size = Pt(9)

    agree, comment = get("agree"), get("comment")
    if agree or comment:
        p = doc.add_paragraph()
        p.add_run(f"{t['reviewer']}: ").bold = True
        a = str(agree) if agree else ""
        p.add_run((AGREE_RU.get(a, a) if lang == "ru" else a) + (f"; {comment}" if comment else ""))


# ---------------------------------------------------------------- tool output (checked_*.xlsx)
# The checking tool (check_descriptions.py) writes a sheet "checked" with one row per
# description: verdicts and findings for items 1, 3, 4(a), 6(a), 6(b).
TOOL_ITEMS = [  # (item, verdict header start, findings header start)
    ("1", "item 1", "findings 1"),
    ("3", "item 3\npartners", "findings 3"),
    ("4(a)", "item 4(a)", "findings 4(a)"),
    ("4(b)", "item 4(b)", "findings 4(b)"),          # version 2 of the tool only
    ("6(a)", "6(a)", "findings 6(a)"),
    ("6(b)", "6(b)", "findings 6(b)"),
]
TOOL_NAMES = {
    "ru": {"1": "Сигналы и число протонов", "3": "Партнёры по КССВ", "4(a)": "Число КССВ",
           "4(b)": "Число партнёров (структура)",
           "6(a)": "Магнитная неэквивалентность", "6(b)": "Сильная связь"},
    "en": {"1": "Signals and proton count", "3": "Coupling partners", "4(a)": "Number of J",
           "4(b)": "Partners in the structure",
           "6(a)": "Magnetic non-equivalence", "6(b)": "Strong coupling"},
}
TOOL_TEXT = {
    "ru": dict(
        title="Что проверяет скрипт",
        rows=[("1. Сигналы и число протонов", "Сумма протонов в описании и число протонов структуры (или формулы)."),
              ("2. Химические сдвиги*", "Сдвиги согласуются с характеристическими диапазонами фрагментов структуры."),
              ("3. Партнёры по КССВ", "У каждой КССВ есть сигнал-партнёр с той же КССВ."),
              ("4(a). Число КССВ", "Число КССВ соответствует мультиплетности (d – одна, dd – две и т. д.)."),
              ("4(b). Число партнёров*", "Мультиплетность не требует больше партнёров по КССВ, чем даёт структура."),
              ("6(a). Магнитная неэквивалентность", "Сигналы колец типа AA′XX′ (пара-дизамещённые бензолы и т. п.), записанные как первого порядка."),
              ("6(b). Сильная связь", "Пары сигналов с общей КССВ и Δν/J < 6 (6–8 – пограничный случай), записанные как первого порядка."),
              ("7. Кажущаяся мультиплетность*", "Мультиплетность, возникающая лишь при случайном равенстве КССВ с неэквивалентными соседями, помечена как кажущаяся (app, J_app) или приведена с отдельными КССВ.")],
        note="* В текущей версии скрипт этот пункт не проверяет; его оценивает человек (или модель) по протоколу. "
             "Скрипт – первичный фильтр: каждая находка – указание, что проверить; окончательное решение "
             "принимает человек по спектру.",
        colours="Цвет ячейки «Вердикт»: красный – несогласовано; жёлтый – проверить, пограничный случай, "
                "нет структуры; зелёный – согласовано. Сигналы в описании: красный фон – несогласованность "
                "по находкам скрипта, жёлтый – проверить.",
        row4b=("4(b). Число партнёров**", "Мультиплетность не требует больше партнёров по КССВ, чем даёт структура."),
        note4b="** Проверяется частично (версия 2): без отнесения скрипт проверяет, есть ли в структуре хоть один "
               "протон (подходящего типа для сдвига), у которого может быть такая мультиплетность; если нет – "
               "«кандидат (проверить)». Мультиплетность, невозможная именно для своего протона, но возможная для "
               "другого, не обнаруживается."),
    "en": dict(
        title="What the script checks",
        rows=[("1. Signals and proton count", "Sum of protons in the description vs protons of the structure (or formula)."),
              ("2. Chemical shifts*", "Shifts agree with the characteristic ranges of the structural fragments."),
              ("3. Coupling partners", "Every J has a partner signal with the same J."),
              ("4(a). Number of J", "The number of J values matches the multiplicity (d - one, dd - two, etc.)."),
              ("4(b). Number of partners*", "The multiplicity needs no more coupling partners than the structure provides."),
              ("6(a). Magnetic non-equivalence", "AA'XX'-type ring signals (para-disubstituted benzenes etc.) reported as first order."),
              ("6(b). Strong coupling", "Pairs of signals with a shared J and Δν/J < 6 (6-8 borderline) reported as first order."),
              ("7. Apparent multiplicity*", "A multiplicity that arises only from accidentally equal couplings to non-equivalent neighbours is marked as apparent (app, J_app) or reported with separate J values.")],
        note="* Not checked by the script in the current version; assessed by a person (or a model) following the "
             "protocol. The script is a first-pass screen: every finding is a pointer to check; the final decision "
             "is made by a person with the spectrum.",
        colours="Verdict cell colour: red - inconsistent; yellow - to check, borderline, no structure; "
                "green - consistent. Signals in the description: red background - inconsistent by the "
                "script's findings, yellow - to check.",
        row4b=("4(b). Number of partners**", "The multiplicity needs no more coupling partners than the structure provides."),
        note4b="** Checked partly (version 2): without an assignment the script checks whether ANY proton of the "
               "structure (of a kind that fits the shift) could have this multiplicity; if none could - "
               "\"candidate (check)\". A multiplicity impossible for its own proton but possible for another one "
               "is not found."),
}
TOOL_VERDICT_RU = [  # (English start, Russian) - longest first
    ("inconsistent? (check)", "возможно несогласовано (проверить)"),
    ("inconsistent (check)", "несогласовано (проверить)"),
    ("inconsistent", "несогласовано"),
    ("partner unverifiable? (check)", "партнёр непроверяем (проверить)"),
    ("partner unverifiable", "партнёр непроверяем"),
    ("borderline (subcategory)", "пограничный случай (подкатегория)"),
    ("candidate (check)", "кандидат (проверить)"),
    ("not checked: mixture/rotamers", "не проверялось: смесь/ротамеры"),
    ("not checked: nuclei not counted", "не проверялось: ядра, которые не учитываются (D, B, металл, NH4+)"),
    ("nothing to check", "нечего проверять"),
    ("no structure", "нет структуры"),
    ("no pair (nothing to check)", "нет пар (нечего проверять)"),
    ("consistent", "согласовано"),
]


def tool_verdict(v, lang):
    if lang == "en":
        return v
    for en, ru in TOOL_VERDICT_RU:
        if v.startswith(en):
            return ru + v[len(en):]
    return v


def tool_shade(v):
    if v.startswith("inconsistent") and "?" not in v:
        return SHADE["inconsistent"]
    if v == "consistent":
        return SHADE["consistent"]
    return SHADE_OTHER


def add_tool_legend(doc, lang, has4b=False):
    tt = TOOL_TEXT[lang]
    doc.add_heading(tt["title"], level=1)
    tab = doc.add_table(rows=0, cols=2)
    tab.style = "Table Grid"
    tab.autofit = False
    rows = [tt["row4b"] if has4b and name.startswith("4(b)") else (name, text) for name, text in tt["rows"]]
    for name, text in rows:
        cells = tab.add_row().cells
        cells[0].text = name
        cells[0].paragraphs[0].runs[0].bold = True
        cells[1].text = text
        if name.endswith("*") and not name.endswith("**"):   # not checked by the script: grey
            for c in cells:
                shade(c, "EDEDED")
                for run in c.paragraphs[0].runs:
                    run.italic = True
    for i, w in enumerate((Cm(6.5), Cm(19.2))):
        tab.columns[i].width = w
        for r in tab.rows:
            r.cells[i].width = w
    doc.add_paragraph()
    doc.add_paragraph(tt["note"])
    if has4b:
        doc.add_paragraph(tt["note4b"])
    doc.add_paragraph(tt["colours"])
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def export_checked(xlsx, out, a):
    """Word file for one output table of the checking tool (sheet 'checked')."""
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    rows = list(wb["checked"].iter_rows(values_only=True))
    header = [str(h) if h is not None else "" for h in rows[0]]
    col = lambda name: find_col(header, name)
    c = {"id": col("id"), "smiles": col("smiles"), "desc": col("description"),
         "note": col("note"), "info": col("info")}
    if c["id"] is None or c["desc"] is None:
        print(f"{Path(xlsx).name}: columns 'id'/'description' not found - skipped")
        return 0
    items = [(it, col(vh), col(fh)) for it, vh, fh in TOOL_ITEMS]
    t, names = TXT[a.lang], TOOL_NAMES[a.lang]

    doc = new_document()
    if a.legend:
        add_tool_legend(doc, a.lang, has4b=col("item 4(b)") is not None)
    data = [r for r in rows[1:] if r[c["id"]] is not None]
    smiles_of = lambda r: r[c["smiles"]] if c["smiles"] is not None else None
    heads = [f"{r[c['id']]} {struct_mark(smiles_of(r), a.lang)}" for r in data]
    if a.contents:
        add_contents(doc, [(h, anchor_name(i)) for i, h in enumerate(heads)], a.lang, a.legend)
    doc.add_heading(Path(xlsx).stem, level=1)
    n = 0
    for i, r in enumerate(data):
        if n:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        n += 1
        h = doc.add_heading(heads[i], level=2)
        add_bookmark(h, anchor_name(i), i)
        smi = smiles_of(r)
        png = structure_png(smi)
        if png is not None:
            doc.add_picture(png, width=Cm(9))
        if smi:
            run = doc.add_paragraph().add_run(str(smi))
            run.font.size, run.font.name = Pt(8), "Consolas"
        p = doc.add_paragraph()
        p.add_run(f"{t['desc']}: ").bold = True
        add_description(p, str(r[c["desc"]] or ""), signal_marks(r, col))

        tab = doc.add_table(rows=1, cols=3)
        tab.style = "Table Grid"
        tab.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate([t["item"], t["verdict"], t["comment"]]):
            tab.rows[0].cells[i].text = h
            tab.rows[0].cells[i].paragraphs[0].runs[0].bold = True
        for it, vi, fi in items:
            if vi is None:
                continue
            v = "" if r[vi] is None else str(r[vi])
            cells = tab.add_row().cells
            cells[0].text = f"{it}. {names[it]}"
            cells[1].text = tool_verdict(v, a.lang)
            cells[2].text = "" if fi is None or r[fi] is None else greek(str(r[fi]))
            if v:
                shade(cells[1], tool_shade(v))
        extra = " ".join(str(r[c[k]]) for k in ("note", "info") if c[k] is not None and r[c[k]])
        if extra:
            cells = tab.add_row().cells
            cells[0].text = t["other"]
            cells[1].merge(cells[2]).text = greek(extra)
        tab.autofit = False
        for i, w in enumerate((Cm(5.2), Cm(5.0), Cm(15.5))):
            tab.columns[i].width = w
            for row in tab.rows:
                row.cells[i].width = w
                for par in row.cells[i].paragraphs:
                    for run in par.runs:
                        run.font.size = Pt(9)
    out = Path(out)
    out.mkdir(parents=True, exist_ok=True)
    path = out / f"{Path(xlsx).stem}.docx"
    doc.save(path)
    print(f"{Path(xlsx).name}: {n} descriptions written to {path.resolve()}")
    return 1


def export_one(xlsx, out, a):
    """Write the Word file(s) for one review workbook into the folder out."""
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    if "assessment" not in wb.sheetnames:
        if "checked" in wb.sheetnames:            # output table of the checking tool
            return export_checked(xlsx, out, a)
        print(f"{Path(xlsx).name}: neither sheet 'assessment' (review) nor 'checked' (tool output) - skipped")
        return 0
    ws = wb["assessment"]
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    c = {k: find_col(header, v) for k, v in COL.items()}
    for item, h in ITEMS:
        c["v" + item] = find_col(header, h)
    missing = [k for k in ("order", "id", "desc") if c[k] is None]
    if missing:
        print(f"{Path(xlsx).name}: columns not found {missing} - skipped")
        return 0

    doi = {}
    if "papers" in wb.sheetnames:
        for r in list(wb["papers"].iter_rows(values_only=True))[1:]:
            if r and r[0] is not None:
                doi[str(r[0])] = r[1]

    groups = OrderedDict()
    for r in rows[1:]:
        if r[c["id"]] is None:
            continue
        groups.setdefault(str(r[c["order"]]), []).append(r)
    if a.articles:
        keep = {s.strip() for s in a.articles.split(",")}
        groups = OrderedDict((k, v) for k, v in groups.items() if k in keep)

    out = Path(out)
    out.mkdir(parents=True, exist_ok=True)
    t = TXT[a.lang]
    written = []
    doc = new_document() if a.one_file else None
    if a.one_file and a.legend:
        add_legend(doc, a.lang)
    if a.one_file and a.contents:
        add_contents(doc, [(review_head(r, c, a.lang), anchor_name(k)) for k, r in
                           enumerate(r for rs in groups.values() for r in rs)], a.lang, a.legend)
    first = True
    k = 0
    for order, rs in groups.items():
        if not a.one_file:
            doc = new_document()
            if a.legend:
                add_legend(doc, a.lang)
            if a.contents:
                add_contents(doc, [(review_head(r, c, a.lang), anchor_name(k + j)) for j, r in enumerate(rs)],
                             a.lang, a.legend)
            first = True
        if not first:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        title = f"{t['article']} {order}" + (f": {doi[order]}" if order in doi else "")
        doc.add_heading(title, level=1)
        for i, r in enumerate(rs):
            add_compound(doc, r, c, a.lang, a.with_script, new_page=(i > 0), anchor=anchor_name(k))
            k += 1
        first = False
        if not a.one_file:
            name = f"{int(order):03d}" if order.isdigit() else safe(order)
            if order in doi:
                name += "_" + safe(doi[order]).replace("10.1021_", "")
            path = out / f"{name}.docx"
            doc.save(path)
            written.append(path)
    if a.one_file:
        path = out / f"review_{a.lang}.docx"
        doc.save(path)
        written.append(path)
    print(f"{Path(xlsx).name}: {len(written)} file(s) written to {out.resolve()}")
    return len(written)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsx", nargs="*",
                    help="review workbook(s); without names, every .xlsx in data/ next to this script")
    ap.add_argument("-o", "--out",
                    help="output folder (default: result/<workbook name>/ next to this script)")
    ap.add_argument("--lang", choices=["ru", "en"], default="ru")
    ap.add_argument("--one-file", action="store_true", help="everything in one .docx")
    ap.add_argument("--with-script", action="store_true", help="add a column with the script's findings")
    ap.add_argument("--articles", help="comma-separated article order numbers")
    ap.add_argument("--no-legend", dest="legend", action="store_false",
                    help="do not add the first page describing the criteria")
    ap.add_argument("--no-contents", dest="contents", action="store_false",
                    help="do not add the list of compounds (contents) after the criteria page")
    a = ap.parse_args()

    here = Path(__file__).resolve().parent
    files = [Path(f) for f in a.xlsx]
    if not files:
        files = sorted(p for p in (here / "data").glob("*.xlsx") if not p.name.startswith("~$"))
        if not files:
            sys.exit(f"No .xlsx files in {here / 'data'}")
    for f in files:
        out = Path(a.out) / f.stem if (a.out and len(files) > 1) else (
            Path(a.out) if a.out else here / "result" / (f.stem + ("_EN" if a.lang == "en" else "")))
        export_one(f, out, a)
    if not HAVE_RDKIT:
        print("rdkit is not installed: Word files have SMILES but no structure images")


if __name__ == "__main__":
    main()
